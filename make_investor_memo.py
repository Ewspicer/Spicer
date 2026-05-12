from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Margins ───────────────────────────────────────────────────────────────
s = doc.sections[0]
s.top_margin    = Inches(0.85)
s.bottom_margin = Inches(0.85)
s.left_margin   = Inches(1.0)
s.right_margin  = Inches(1.0)

# ── Helpers ───────────────────────────────────────────────────────────────
NAVY   = RGBColor(15, 60, 110)
RED    = RGBColor(180, 30, 30)
GOLD   = RGBColor(150, 100, 0)
GRAY   = RGBColor(100, 100, 100)
LGRAY  = RGBColor(160, 160, 160)

def rule(doc, color='CCCCCC'):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single')
    bot.set(qn('w:sz'),  '4')
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), color)
    pBdr.append(bot)
    pPr.append(pBdr)

def heading(doc, label, title, size=13):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(3)
    r1 = p.add_run(label + '  ')
    r1.bold = True; r1.font.size = Pt(9); r1.font.color.rgb = GOLD
    r2 = p.add_run(title)
    r2.bold = True; r2.font.size = Pt(size); r2.font.color.rgb = NAVY

def body(doc, text, size=10.5, space_after=7, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p

def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(3)
    r = p.add_run(text)
    r.font.size = Pt(size)

# ══════════════════════════════════════════════════════════════════════════
# HEADER
# ══════════════════════════════════════════════════════════════════════════
conf = doc.add_paragraph()
conf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
conf.paragraph_format.space_before = Pt(0)
conf.paragraph_format.space_after  = Pt(3)
r = conf.add_run('CONFIDENTIAL — BOARD & INVESTOR RELATIONS USE ONLY')
r.font.size = Pt(7.5); r.font.color.rgb = LGRAY

title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after  = Pt(3)
r = title_p.add_run('INVESTOR DAY POSITIONING MEMO')
r.bold = True; r.font.size = Pt(20); r.font.color.rgb = NAVY

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.paragraph_format.space_before = Pt(0)
sub_p.paragraph_format.space_after  = Pt(2)
r = sub_p.add_run('Meridian Technologies (MRDN)  |  March 11, 2026  |  Prepared by: Catherine Park, CEO')
r.font.size = Pt(9); r.font.color.rgb = GRAY

rule(doc, '1a6eb5')

# ══════════════════════════════════════════════════════════════════════════
# EXEC SUMMARY
# ══════════════════════════════════════════════════════════════════════════
heading(doc, 'EXECUTIVE SUMMARY', '')

es = doc.add_paragraph()
es.paragraph_format.space_before = Pt(0)
es.paragraph_format.space_after  = Pt(8)
r = es.add_run(
    'Meridian Technologies is declaring itself the '
)
r.font.size = Pt(10.5)
r2 = es.add_run('agentic work platform that enterprise governance teams trust')
r2.bold = True; r2.font.size = Pt(10.5); r2.font.color.rgb = NAVY
r3 = es.add_run(
    ' — built on the most auditable, compliance-ready agent infrastructure in the '
    'collaboration software category. This memo records the strategic rationale for that '
    'declaration, the competitive evidence that supports it, the risks we are carrying '
    'knowingly, and the three commitments we are making to investors today. '
    'Meridian enters Investor Day with $506M in liquidity, zero debt, a 44% enterprise '
    'Copilot attach rate, and a governance moat that no direct competitor has matched. '
    'We are moving — deliberately and with financial capacity — from the center of the '
    'competitive map to the upper-right quadrant: agentic ambition, premium pricing, '
    'enterprise trust. That quadrant has one occupant today (Atlassian, in developer '
    'tooling). Meridian will own it in project management and regulated-industry work.'
)
r3.font.size = Pt(10.5)

rule(doc)

# ══════════════════════════════════════════════════════════════════════════
# POSITION
# ══════════════════════════════════════════════════════════════════════════
heading(doc, 'SECTION 1', 'The Position')

body(doc,
    'Effective today, Meridian\'s public positioning is:',
    space_after=4)

quote_p = doc.add_paragraph()
quote_p.paragraph_format.space_before = Pt(4)
quote_p.paragraph_format.space_after  = Pt(4)
quote_p.paragraph_format.left_indent  = Inches(0.35)
quote_p.paragraph_format.right_indent = Inches(0.35)
r = quote_p.add_run(
    '"Meridian is the agentic work platform for enterprises that cannot afford to get AI wrong — '
    'built on the governance infrastructure that regulated industries already trust."'
)
r.bold = True; r.italic = True; r.font.size = Pt(11); r.font.color.rgb = NAVY

body(doc,
    'This replaces "the project management and team collaboration platform." '
    'PM is not abandoned — it is repositioned as the surface through which agents operate, '
    'not the product we are selling. The pricing model will shift in H2 2026 to a '
    'per-seat-plus-consumption hybrid. The sales motion will shift to agent-led for new '
    'logos and net-new expansion; the PM motion is preserved for renewals.',
    space_after=8)

# ══════════════════════════════════════════════════════════════════════════
# WHY
# ══════════════════════════════════════════════════════════════════════════
heading(doc, 'SECTION 2', 'Why This Position — The Strategic Logic')

body(doc,
    'Four inputs drove this decision:',
    space_after=3)

bullet(doc,
    'Customer signal: In 18 of 22 enterprise advisory board sessions (Dec 2025–Feb 2026), '
    'the top ask was agent governance — auditability, role-based agent permissions, model '
    'selection controls, data residency. No competitor has built all four. Meridian is closest.')
bullet(doc,
    'Competitive white space: Asana and Atlassian have both claimed "agentic platform" — '
    'but neither has the governance depth for regulated industries. Smartsheet has the '
    'governance posture but has explicitly rejected the word "agentic." The upper-right '
    'quadrant of the competitive map (agentic + premium + governed) is unclaimed in PM.')
bullet(doc,
    'Financial capacity: $506M in liquidity, $67M in 2025 FCF, $250M in estimated M&A '
    'capacity, and $55M in AI-specific R&D already budgeted. We can absorb an 18-month '
    'revenue ramp lag without covenant risk or capital raises.')
bullet(doc,
    'Team signal: The Helio acquisition gave us an agent framework and 28 applied AI '
    'engineers. They did not join to build AI features for a PM tool. Option B retains '
    'the talent that makes the roadmap executable.')

body(doc, '', space_after=2)
rule(doc)

# ══════════════════════════════════════════════════════════════════════════
# COMPETITIVE LANDSCAPE TABLE
# ══════════════════════════════════════════════════════════════════════════
heading(doc, 'SECTION 3', 'Competitive Landscape')

body(doc,
    'The table below maps the four primary competitors across the dimensions most relevant '
    'to Meridian\'s positioning decision. Sources: competitors_cached/ (Feb 2026 snapshots).',
    space_after=5)

# Table
cols = ['Competitor', 'AI Positioning', 'Pricing Posture',
        'Latest Flagship Announcement', 'Closer to Option…']
rows = [
    ['Asana',
     '"Work management platform with AI built in." Calls itself an "agent operating system." De-emphasizes classic PM in marketing.',
     'Bundled into Advanced+ tiers at no extra cost. No consumption pricing. Explicitly rejected consumption model.',
     'Nov 2025: AI Studio and Smart Workflows bundled into Advanced/Enterprise — direct competitive escalation vs. add-on competitors.',
     'B — agentic claim, no governance depth'],
    ['Monday.com',
     '"Work OS supercharged with AI everywhere." AI as a horizontal layer, not a separate product. Breadth and price as differentiators.',
     'Bundled into Pro+ at no extra cost. No consumption pricing. "Consumption creates buying friction we don\'t want."',
     'Jan 2026: monday AI Agents GA — customer-buildable agents in natural language, bundled into Pro tier.',
     'A/B hybrid — agentic framing, bundled execution; no vertical or governance moat'],
    ['Smartsheet',
     '"Enterprise work platform you can trust with AI." Avoids "agentic." Emphasizes trust, compliance, audit-committee approval.',
     'AI bundled into Business/Enterprise tiers. AI Compliance Pack as premium add-on (~$15/seat/month).',
     'Dec 2025: AI Compliance Pack — agent-level audit logs, model selection, data residency. Priced as enterprise premium.',
     'A — PM-centric, governance-forward, cautious on agentic ambition'],
    ['Atlassian',
     '"Agentic enterprise platform for software and IT teams." Most explicit agentic bet; internal language: "an agent company that ships software."',
     'Rovo sold as separate paid product: per-seat (~$20/seat/mo) + per-action consumption. Only competitor with consumption pricing today.',
     'Jan 2026: Rovo Studio — developer-targeted agent builder on Jira/Confluence/Bitbucket infrastructure. Consumption-priced.',
     'B — closest to Meridian\'s Option B, but dev/IT-shaped not PM-shaped; thin outside engineering orgs'],
]

table = doc.add_table(rows=1 + len(rows), cols=len(cols))
table.style = 'Table Grid'

# Header row
hdr = table.rows[0]
hdr_fills = ['1a4f82'] * len(cols)
for i, (cell, col) in enumerate(zip(hdr.cells, cols)):
    cell.text = col
    run = cell.paragraphs[0].runs[0]
    run.bold = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(255, 255, 255)
    # Cell shading
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), '1a4f82')
    tcPr.append(shd)

col_widths = [Inches(0.9), Inches(1.55), Inches(1.35), Inches(1.8), Inches(1.0)]

for i, row_data in enumerate(rows):
    row = table.rows[i + 1]
    fill = 'f2f6fc' if i % 2 == 0 else 'ffffff'
    for j, (cell, text) in enumerate(zip(row.cells, row_data)):
        cell.width = col_widths[j]
        cell.text = text
        p = cell.paragraphs[0]
        run = p.runs[0] if p.runs else p.add_run(text)
        run.font.size = Pt(8)
        if j == 0:
            run.bold = True
        if j == len(cols) - 1:
            run.font.color.rgb = NAVY
            run.bold = True
        # Row shading
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), fill)
        tcPr.append(shd)

doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ── Matrix chart ──────────────────────────────────────────────────────────
body(doc,
    'Figure 1: Competitive Positioning Matrix — AI Strategy × Pricing Posture',
    size=9, space_after=3)
body(doc,
    'The matrix below plots all four competitors and Meridian\'s current and proposed '
    'positions. The target white space (upper right: agentic + premium governance) '
    'has one occupant (Atlassian, in developer tooling). Meridian\'s proposed move '
    'positions it as the agentic+governance platform for PM and regulated industries.',
    size=9, space_after=4)

img_p = doc.add_paragraph()
img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
img_p.paragraph_format.space_before = Pt(0)
img_p.paragraph_format.space_after  = Pt(6)
img_p.add_run().add_picture('/home/user/Spicer/meridian_competitive_matrix.png', width=Inches(5.6))

rule(doc)

# ══════════════════════════════════════════════════════════════════════════
# THREE RISKS
# ══════════════════════════════════════════════════════════════════════════
heading(doc, 'SECTION 4', 'Risks We Are Carrying Knowingly')

risks = [
    ('Risk 1 — The revenue air gap.',
     'Declaring "agentic platform" and transitioning to consumption pricing while '
     'deemphasizing the PM sales motion creates a 12–18 month period in which '
     'agentic ARR must ramp before PM growth slows further. AI Copilot ended 2025 at '
     '$3.5M ARR. If enterprise consumption adoption is slower than modeled — customers '
     'told us they want human-in-the-loop checkpoints, not autonomous workflows — '
     '2026 ARR growth could land at the low end of guidance (10%) or below. '
     'Mitigation: $506M liquidity buffer; FCF positive at any growth rate above 5%; '
     'PM renewals preserved under both motions.'),
    ('Risk 2 — Atlassian moves into our verticals.',
     'Atlassian is the only competitor already in the upper-right quadrant of the '
     'positioning map, and their install base is enormous. If they decide to invest '
     'in financial services and life sciences vertical depth — or acquire a compliance '
     'tooling company — our differentiated moat narrows. '
     'Mitigation: our FedRAMP Moderate, HIPAA, and (2026) GxP certifications are '
     '12–18 months ahead; six of the top-10 North American retail banks are active '
     'customers; Atlassian\'s developer-shaped product is a structural disadvantage '
     'in these verticals.'),
    ('Risk 3 — Helio retention and talent execution.',
     'The agent builder roadmap depends on retaining the Helio team through their '
     '2026 compensation cliff. Twenty-six of twenty-eight Helio engineers remain '
     'today. The compensation program runs through 2029, but the largest vesting event '
     'is in mid-2026. If key engineers leave after that cliff, the Q2 2026 '
     'agent-builder shipment is at risk and the Option B narrative loses its '
     'primary technical proof point. '
     'Mitigation: People & Compensation Committee to review and top-up retention '
     'agreements by Q1 2026; broader senior IC engineering refresh approved concurrently.'),
]

for title, text in risks:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(2)
    r1 = p.add_run(title + '  ')
    r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = RED
    body(doc, text, size=10, space_after=5)

rule(doc)

# ══════════════════════════════════════════════════════════════════════════
# THREE COMMITMENTS
# ══════════════════════════════════════════════════════════════════════════
heading(doc, 'SECTION 5', 'Three Commitments to Investors')

commitments = [
    ('Commitment 1 — Ship the governance suite in Q1 2026.',
     'The enterprise governance suite for AI agents (audit logs, role-based agent '
     'permissions, model selection, data residency controls) will reach GA in Q1 2026. '
     'This is the product that converts our positioning claim into a proof point. '
     'We will report enterprise governance suite attach rate at each quarterly call '
     'starting Q1 2026.'),
    ('Commitment 2 — Announce consumption pricing model in H2 2026.',
     'We will announce the hybrid per-seat-plus-consumption pricing model no later '
     'than Q3 2026. We will provide investors with a framework for modeling '
     'consumption revenue contribution — including assumptions, leading indicators, '
     'and a range of outcomes — at the time of announcement. We will not guide '
     'consumption revenue until we have two full quarters of data.'),
    ('Commitment 3 — Report three new Copilot metrics beginning Q1 2026.',
     'As committed on the Q3 2025 earnings call: (1) Copilot attach rate on enterprise '
     'renewals (target: above 40% sustained; Q4 2025 actual: 44%), '
     '(2) weekly agentic-feature usage rate across active users, and '
     '(3) Copilot appearance in win-loss data. These are the metrics by which '
     'we ask investors to judge our AI execution.'),
]

for title, text in commitments:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(2)
    r1 = p.add_run(title + '  ')
    r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = NAVY
    body(doc, text, size=10, space_after=5)

rule(doc)

# ── Footer ────────────────────────────────────────────────────────────────
foot = doc.add_paragraph()
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
foot.paragraph_format.space_before = Pt(6)
r = foot.add_run(
    'Sources: meridian_internal_brief.md · meridian_ai_strategy_options.md · '
    'meridian_recent_customer_feedback.md · meridian_financials_summary.csv · '
    'competitors_cached/ (Feb 2026)  |  '
    'For board and IR use only. Not for public distribution prior to March 11, 2026.'
)
r.font.size = Pt(7.5); r.font.color.rgb = LGRAY

doc.save('/home/user/Spicer/investor_day_positioning_memo.docx')
print("Memo saved.")
